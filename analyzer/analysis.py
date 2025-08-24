import pandas as pd
import statsmodels.api as sm
from scipy import stats
from docx import Document
import matplotlib.pyplot as plt
import seaborn as sns
import os
import uuid
import uuid
import os
import seaborn as sns
import matplotlib.pyplot as plt
import statsmodels.api as sm
from docx import Document
from docx.shared import Inches
# --- Step 1: Dataset Description ---
import io
import sys
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.shared import Inches

def perform_eda(df):
    summary_html = df.describe().to_html(classes="table")
    doc = Document()
    doc.add_heading("Exploratory Data Analysis", level=1)
    doc.add_paragraph(df.describe().to_string())

    plots = []
    for col in df.select_dtypes(include=['int', 'float']):
        plt.figure()
        sns.histplot(df[col], kde=True)
        plot_name = f"static/plots/{uuid.uuid4().hex}_{col}.png"
        plt.title(f"Histogram - {col}")
        plt.savefig(plot_name)
        plots.append(plot_name)
        doc.add_picture(plot_name, width=docx.shared.Inches(5))
        plt.close()

    return summary_html, doc, plots



from docx.shared import Inches
import matplotlib.pyplot as plt
import seaborn as sns
import statsmodels.api as sm
from docx import Document
import os
import uuid

# def preprocess_data(df):
#     steps = []

#     # Drop duplicates
#     before = df.shape[0]
#     df = df.drop_duplicates()
#     after = df.shape[0]
#     if before != after:
#         steps.append(f"Removed {before - after} duplicate rows.")

#     # Handle missing values
#     missing = df.isnull().sum().sum()
#     if missing > 0:
#         df = df.dropna()  # or fillna if preferred
#         steps.append(f"Removed {missing} missing values by dropping rows.")

#     # Convert categorical if needed (optional)
#     # for col in df.select_dtypes(include='object'):
#     #     df[col] = df[col].astype('category')

#     return df, "\n".join(steps)



def simple_regression(df, x_col, y_col):
    # --- Step 1: Dataset Description ---
    num_rows, num_columns = df.shape

    # Count numerical and categorical columns
    numerical_cols = df.select_dtypes(include=['number']).columns
    categorical_cols = df.select_dtypes(include=['object', 'category']).columns
    num_numerical = len(numerical_cols)
    num_categorical = len(categorical_cols)

    # Dataset overview string
    dataset_info = (
        f"The dataset contains <strong>{num_rows}</strong> rows and <strong>{num_columns}</strong> columns.<br>"
        f"It includes <strong>{num_numerical}</strong> numerical column(s) and "
        f"<strong>{num_categorical}</strong> categorical column(s).<br><br>"
    )

    dataset_info_doc = (
        f"The dataset contains {num_rows}rows and {num_columns} columns.<br>"
        f"It includes {num_numerical} numerical column(s) and "
        f"{num_categorical} categorical column(s)."
    )

    # --- Capture df.info() as string ---
    buffer = io.StringIO()
    df.info(buf=buffer)
    info_str = buffer.getvalue()
    info_html = f"<pre>{info_str}</pre>"

    # --- Describe table as HTML ---
    describe_html = df.describe(include='all').to_html(classes='table table-striped', border=0)

    # --- Final Dataset Description Block ---
    dataset_description = dataset_info + info_html + "<br>" + describe_html

    # --- Step 2: Preprocessing Steps ---
    preprocessing_steps = []

    # Drop duplicates
    initial_rows = df.shape[0]
    df = df.drop_duplicates()
    final_rows = df.shape[0]
    if initial_rows != final_rows:
        preprocessing_steps.append(f"Removed {initial_rows - final_rows} duplicate rows.")

    # Drop missing values
    missing_values = df.isnull().sum().sum()
    if missing_values > 0:
        df = df.dropna()
        preprocessing_steps.append(f"Dropped rows with {missing_values} missing values.")

    if not preprocessing_steps:
        preprocessing_steps.append("No preprocessing was required.")

    # Preprocessing steps formatted for HTML
    preprocessing_steps_text = "<br>".join(preprocessing_steps)



    # --- Step 3: Regression Analysis ---
    X = sm.add_constant(df[x_col])
    y = df[y_col]
    model = sm.OLS(y, X).fit()

    summary_html = model.summary().tables[1].as_html()
    objective = f"Linear regression to predict {y_col} from {x_col}."
    interpretation = (
        f"The coefficient for {x_col} is {round(model.params[x_col], 4)} with a p-value of "
        f"{round(model.pvalues[x_col], 4)}. "
        "This indicates a statistically significant relationship."
        if model.pvalues[x_col] < 0.05
        else f"The coefficient for {x_col} is {round(model.params[x_col], 4)} but is not statistically significant."
    )

    # --- Step 4: Create Plots ---
    plot_filenames = []
    os.makedirs("static/plots", exist_ok=True)

    # Scatter plot with regression line
    plt.figure(figsize=(6, 4))
    sns.regplot(x=df[x_col], y=df[y_col])
    plt.title(f"{y_col} vs {x_col}")
    scatter_plot_name = f"{uuid.uuid4().hex}_regression_plot.png"
    plt.savefig(f"static/plots/{scatter_plot_name}", bbox_inches='tight')
    plt.close()
    plot_filenames.append(scatter_plot_name)

    # Residual plot
    plt.figure(figsize=(6, 4))
    sns.residplot(x=model.fittedvalues, y=model.resid, lowess=True)
    plt.xlabel("Fitted Values")
    plt.ylabel("Residuals")
    plt.title("Residuals vs Fitted")
    resid_plot_name = f"{uuid.uuid4().hex}_resid_plot.png"
    plt.savefig(f"static/plots/{resid_plot_name}", bbox_inches='tight')
    plt.close()
    plot_filenames.append(resid_plot_name)

    # --- Step 5: Word Report ---
    doc = Document()
    # Assuming 'doc' is your Document() object

    # --- Add Dataset Structure (df.info()) ---
    doc.add_heading("Structure and Data Types", level=2)

    # Add info_str in monospaced font
    p = doc.add_paragraph()
    run = p.add_run(info_str)
    run.font.name = 'Courier New'
    run.font.size = Pt(10)

    # Ensure style is monospaced on all platforms
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), 'Courier New')
    doc.add_heading("Dataset Description", level=2)
    doc.add_paragraph(dataset_info_doc)
    

    doc.add_heading("Preprocessing Steps", level=2)
    doc.add_paragraph(preprocessing_steps_text)

    doc.add_heading("Simple Linear Regression", level=1)
    doc.add_heading("Objective", level=2)
    doc.add_paragraph(objective)

    doc.add_heading("Regression Results", level=2)
    doc.add_paragraph(str(model.summary()))

    doc.add_heading("Interpretation", level=2)
    doc.add_paragraph(interpretation)

    doc.add_heading("Visualizations", level=2)
    for plot in plot_filenames:
        doc.add_picture(f"static/plots/{plot}", width=Inches(5))
        doc.add_paragraph("")

    # --- Return everything ---
   # --- Return everything ---
    return {
        "summary_html": summary_html,
        "objective": objective,
        "interpretation": interpretation,
        "plots": plot_filenames,
        "doc": doc,
        "dataset_description": dataset_description,
        "preprocessing_steps": preprocessing_steps_text
    }




def multiple_regression(df, x_cols, y_col):
    # --- Step 1: Dataset Description ---
    num_rows, num_columns = df.shape

    numerical_cols = df.select_dtypes(include=['number']).columns
    categorical_cols = df.select_dtypes(include=['object', 'category']).columns
    num_numerical = len(numerical_cols)
    num_categorical = len(categorical_cols)

    dataset_info = (
        f"The dataset contains <strong>{num_rows}</strong> rows and <strong>{num_columns}</strong> columns.<br>"
        f"It includes <strong>{num_numerical}</strong> numerical column(s) and "
        f"<strong>{num_categorical}</strong> categorical column(s).<br><br>"
    )

    dataset_info_doc = (
        f"The dataset contains {num_rows} rows and {num_columns} columns. "
        f"It includes {num_numerical} numerical column(s) and "
        f"{num_categorical} categorical column(s)."
    )

    # Capture df.info() as string
    buffer = io.StringIO()
    df.info(buf=buffer)
    info_str = buffer.getvalue()
    info_html = f"<pre>{info_str}</pre>"

    describe_html = df.describe(include='all').to_html(classes='table table-striped', border=0)

    dataset_description = dataset_info + info_html + "<br>" + describe_html

    # --- Step 2: Preprocessing Steps ---
    preprocessing_steps = []
    initial_rows = df.shape[0]
    df = df.drop_duplicates()
    final_rows = df.shape[0]
    if initial_rows != final_rows:
        preprocessing_steps.append(f"Removed {initial_rows - final_rows} duplicate rows.")

    missing_values = df.isnull().sum().sum()
    if missing_values > 0:
        df = df.dropna()
        preprocessing_steps.append(f"Dropped rows with {missing_values} missing values.")

    if not preprocessing_steps:
        preprocessing_steps.append("No preprocessing was required.")

    preprocessing_steps_text = "<br>".join(preprocessing_steps)

    # --- Step 3: Multiple Regression ---
    X = sm.add_constant(df[x_cols])   # independent variables
    y = df[y_col]
    model = sm.OLS(y, X).fit()

    summary_html = model.summary().tables[1].as_html()
    objective = f"Multiple linear regression to predict {y_col} from {', '.join(x_cols)}."
    interpretation = (
        "Regression results show how each independent variable contributes "
        "to predicting the dependent variable. "
        "Significance of predictors is determined by their p-values."
    )

    # --- Step 4: Plots ---
    plot_filenames = []
    os.makedirs("static/plots", exist_ok=True)

    # Residuals vs Fitted
    plt.figure(figsize=(6, 4))
    sns.residplot(x=model.fittedvalues, y=model.resid, lowess=True, line_kws={'color': 'red'})
    plt.xlabel("Fitted Values")
    plt.ylabel("Residuals")
    plt.title("Residuals vs Fitted")
    resid_plot_name = f"{uuid.uuid4().hex}_resid_plot.png"
    plt.savefig(f"static/plots/{resid_plot_name}", bbox_inches='tight')
    plt.close()
    plot_filenames.append(resid_plot_name)

    # Q-Q Plot
    plt.figure(figsize=(6, 4))
    sm.qqplot(model.resid, line='45', fit=True)
    plt.title("Normal Q-Q Plot")
    qq_plot_name = f"{uuid.uuid4().hex}_qq_plot.png"
    plt.savefig(f"static/plots/{qq_plot_name}", bbox_inches='tight')
    plt.close()
    plot_filenames.append(qq_plot_name)

    # Actual vs Predicted
    plt.figure(figsize=(6, 4))
    plt.scatter(model.fittedvalues, y, alpha=0.7)
    plt.plot([y.min(), y.max()], [y.min(), y.max()], 'r--')
    plt.xlabel("Predicted Values")
    plt.ylabel("Actual Values")
    plt.title("Actual vs Predicted")
    actual_pred_plot = f"{uuid.uuid4().hex}_actual_vs_pred.png"
    plt.savefig(f"static/plots/{actual_pred_plot}", bbox_inches='tight')
    plt.close()
    plot_filenames.append(actual_pred_plot)

    # --- Step 5: Word Report ---
    doc = Document()

    doc.add_heading("Structure and Data Types", level=2)
    p = doc.add_paragraph()
    run = p.add_run(info_str)
    run.font.name = 'Courier New'
    run.font.size = Pt(10)
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), 'Courier New')

    doc.add_heading("Dataset Description", level=2)
    doc.add_paragraph(dataset_info_doc)

    doc.add_heading("Preprocessing Steps", level=2)
    doc.add_paragraph(preprocessing_steps_text)

    doc.add_heading("Multiple Linear Regression", level=1)
    doc.add_heading("Objective", level=2)
    doc.add_paragraph(objective)

    doc.add_heading("Regression Results", level=2)
    doc.add_paragraph(str(model.summary()))

    doc.add_heading("Interpretation", level=2)
    doc.add_paragraph(interpretation)

    doc.add_heading("Visualizations", level=2)
    for plot in plot_filenames:
        doc.add_picture(f"static/plots/{plot}", width=Inches(5))
        doc.add_paragraph("")

    # --- Return everything ---
    return {
        "summary_html": summary_html,
        "objective": objective,
        "interpretation": interpretation,
        "plots": plot_filenames,
        "doc": doc,
        "dataset_description": dataset_description,
        "preprocessing_steps": preprocessing_steps_text
    }
def one_sample_ttest(df, sample_col, popmean):
    # --- Step 1: Dataset Description ---
    num_rows, num_columns = df.shape

    # Count numerical and categorical columns
    numerical_cols = df.select_dtypes(include=['number']).columns
    categorical_cols = df.select_dtypes(include=['object', 'category']).columns
    num_numerical = len(numerical_cols)
    num_categorical = len(categorical_cols)

    dataset_info = (
        f"The dataset contains <strong>{num_rows}</strong> rows and <strong>{num_columns}</strong> columns.<br>"
        f"It includes <strong>{num_numerical}</strong> numerical column(s) and "
        f"<strong>{num_categorical}</strong> categorical column(s).<br><br>"
    )

    dataset_info_doc = (
        f"The dataset contains {num_rows} rows and {num_columns} columns.\n"
        f"It includes {num_numerical} numerical column(s) and "
        f"{num_categorical} categorical column(s)."
    )

    # --- Capture df.info() as string ---
    buffer = io.StringIO()
    df.info(buf=buffer)
    info_str = buffer.getvalue()
    info_html = f"<pre>{info_str}</pre>"

    # --- Describe table as HTML ---
    describe_html = df.describe(include='all').to_html(classes='table table-striped', border=0)

    # --- Final Dataset Description Block ---
    dataset_description = dataset_info + info_html + "<br>" + describe_html

    # --- Step 2: Preprocessing Steps ---
    preprocessing_steps = []

    initial_rows = df.shape[0]
    df = df.drop_duplicates()
    final_rows = df.shape[0]
    if initial_rows != final_rows:
        preprocessing_steps.append(f"Removed {initial_rows - final_rows} duplicate rows.")

    missing_values = df.isnull().sum().sum()
    if missing_values > 0:
        df = df.dropna()
        preprocessing_steps.append(f"Dropped rows with {missing_values} missing values.")

    if not preprocessing_steps:
        preprocessing_steps.append("No preprocessing was required.")

    preprocessing_steps_text = "<br>".join(preprocessing_steps)

    # --- Step 3: One Sample T-Test ---
    data = df[sample_col].dropna()
    t_stat, p_val = stats.ttest_1samp(data, popmean)

    objective = f"One-sample t-test to check if the mean of '{sample_col}' differs from {popmean}."
    interpretation = (
        f"The sample mean is {data.mean():.4f}. "
        f"T-statistic = {t_stat:.4f}, P-value = {p_val:.4f}. "
        "This indicates a statistically significant difference from the population mean."
        if p_val < 0.05
        else f"The sample mean is {data.mean():.4f}. "
             f"T-statistic = {t_stat:.4f}, P-value = {p_val:.4f}. "
             "This indicates no significant difference from the population mean."
    )

    # --- Step 4: Create Plots ---
    plot_filenames = []
    os.makedirs("static/plots", exist_ok=True)

    # Histogram of sample data with population mean line
    plt.figure(figsize=(6, 4))
    sns.histplot(data, kde=True, bins=15, color="skyblue")
    plt.axvline(popmean, color="red", linestyle="--", label=f"Population Mean ({popmean})")
    plt.axvline(data.mean(), color="green", linestyle="-", label=f"Sample Mean ({data.mean():.2f})")
    plt.title(f"Distribution of {sample_col}")
    plt.legend()
    hist_plot_name = f"{uuid.uuid4().hex}_ttest_hist.png"
    plt.savefig(f"static/plots/{hist_plot_name}", bbox_inches='tight')
    plt.close()
    plot_filenames.append(hist_plot_name)

    # --- Step 5: Word Report ---
    doc = Document()

    doc.add_heading("Structure and Data Types", level=2)
    p = doc.add_paragraph()
    run = p.add_run(info_str)
    run.font.name = 'Courier New'
    run.font.size = Pt(10)
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), 'Courier New')

    doc.add_heading("Dataset Description", level=2)
    doc.add_paragraph(dataset_info_doc)

    doc.add_heading("Preprocessing Steps", level=2)
    doc.add_paragraph(preprocessing_steps_text)

    doc.add_heading("One-Sample T-Test", level=1)
    doc.add_heading("Objective", level=2)
    doc.add_paragraph(objective)

    doc.add_heading("T-Test Results", level=2)
    doc.add_paragraph(f"T-statistic: {t_stat:.4f}\nP-value: {p_val:.4f}\nSample Mean: {data.mean():.4f}")

    doc.add_heading("Interpretation", level=2)
    doc.add_paragraph(interpretation)

    doc.add_heading("Visualization", level=2)
    for plot in plot_filenames:
        doc.add_picture(f"static/plots/{plot}", width=Inches(5))
        doc.add_paragraph("")

    # --- Return everything ---
    return {
        "t_stat": t_stat,
        "p_val": p_val,
        "objective": objective,
        "interpretation": interpretation,
        "plots": plot_filenames,
        "doc": doc,
        "dataset_description": dataset_description,
        "preprocessing_steps": preprocessing_steps_text
    }
def hypothesis_testing(df, col):
    t_stat, p_val = stats.ttest_1samp(df[col], popmean=0)
    result_html = f"<p>T-statistic: {t_stat:.4f}</p><p>P-value: {p_val:.4f}</p>"
    doc = Document()
    doc.add_heading("One-Sample T-Test", level=1)
    doc.add_paragraph(f"T-statistic: {t_stat:.4f}, P-value: {p_val:.4f}")
    return result_html, doc
